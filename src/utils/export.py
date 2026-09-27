"""
Document Export Utilities for Transcription Results
Supports multiple formats: TXT, SRT, DOCX, PDF, JSON
"""
import json
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional
import os


def format_timestamp_srt(seconds: float) -> str:
    """Convert seconds to SRT timestamp format (00:00:00,000)"""
    if seconds is None:
        seconds = 0.0
    
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    milliseconds = int((seconds - int(seconds)) * 1000)
    
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{milliseconds:03d}"


def export_txt(text: str, output_path: str, metadata: Optional[Dict] = None) -> str:
    """
    Export transcription as plain text file
    
    Args:
        text: Transcription text
        output_path: Output file path
        metadata: Optional metadata to include
    
    Returns:
        Path to created file
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        if metadata:
            f.write("=" * 60 + "\n")
            f.write("TRANSCRIPTION METADATA\n")
            f.write("=" * 60 + "\n\n")
            
            if 'filename' in metadata:
                f.write(f"File: {metadata['filename']}\n")
            if 'duration' in metadata:
                f.write(f"Duration: {metadata['duration']}s\n")
            if 'date' in metadata:
                f.write(f"Date: {metadata['date']}\n")
            if 'model' in metadata:
                f.write(f"Model: {metadata['model']}\n")
            if 'word_count' in metadata:
                f.write(f"Words: {metadata['word_count']}\n")
            
            f.write("\n" + "=" * 60 + "\n\n")
        
        f.write(text)
    
    return output_path


def export_srt(segments: List[Dict], output_path: str) -> str:
    """
    Export subtitles as SRT format
    
    Args:
        segments: List of subtitle segments with start, end, text
        output_path: Output file path
    
    Returns:
        Path to created file
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        for i, segment in enumerate(segments, start=1):
            start = segment.get('start', 0)
            end = segment.get('end', start + 2)
            text = segment.get('text', '').strip()
            
            if not text:
                continue
            
            f.write(f"{i}\n")
            f.write(f"{format_timestamp_srt(start)} --> {format_timestamp_srt(end)}\n")
            f.write(f"{text}\n")
            f.write("\n")
    
    return output_path


def export_vtt(segments: List[Dict], output_path: str) -> str:
    """
    Export subtitles as WebVTT format
    
    Args:
        segments: List of subtitle segments
        output_path: Output file path
    
    Returns:
        Path to created file
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("WEBVTT\n\n")
        
        for i, segment in enumerate(segments, start=1):
            start = segment.get('start', 0)
            end = segment.get('end', start + 2)
            text = segment.get('text', '').strip()
            
            if not text:
                continue
            
            # WebVTT uses dots instead of commas for milliseconds
            start_vtt = format_timestamp_srt(start).replace(',', '.')
            end_vtt = format_timestamp_srt(end).replace(',', '.')
            
            f.write(f"{i}\n")
            f.write(f"{start_vtt} --> {end_vtt}\n")
            f.write(f"{text}\n")
            f.write("\n")
    
    return output_path


def export_docx(
    text: str,
    output_path: str,
    title: str = "Transcription",
    metadata: Optional[Dict] = None,
    segments: Optional[List[Dict]] = None
) -> str:
    """
    Export transcription as DOCX (Microsoft Word) file
    
    Args:
        text: Transcription text
        output_path: Output file path
        title: Document title
        metadata: Optional metadata
        segments: Optional segments with timestamps
    
    Returns:
        Path to created file
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError("python-docx not installed. Install with: pip install python-docx")
    
    doc = Document()
    
    # Title
    title_para = doc.add_heading(title, level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata section
    if metadata:
        doc.add_heading('Metadata', level=2)
        meta_table = doc.add_table(rows=0, cols=2)
        meta_table.style = 'Light Grid Accent 1'
        
        for key, value in metadata.items():
            row = meta_table.add_row()
            row.cells[0].text = key.replace('_', ' ').title()
            row.cells[1].text = str(value)
        
        doc.add_paragraph()
    
    # Transcription section
    doc.add_heading('Transcription', level=2)
    
    if segments:
        # If segments available, show with timestamps
        for segment in segments:
            start = segment.get('start', 0)
            end = segment.get('end', 0)
            segment_text = segment.get('text', '').strip()
            
            if not segment_text:
                continue
            
            # Timestamp
            timestamp_para = doc.add_paragraph()
            timestamp_run = timestamp_para.add_run(
                f"[{format_timestamp_srt(start).split(',')[0]} - {format_timestamp_srt(end).split(',')[0]}]"
            )
            timestamp_run.font.size = Pt(9)
            timestamp_run.font.color.rgb = RGBColor(128, 128, 128)
            
            # Text
            text_para = doc.add_paragraph(segment_text)
            text_para.paragraph_format.space_after = Pt(12)
    else:
        # Plain text
        doc.add_paragraph(text)
    
    # Footer
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_run = footer_para.add_run(
        f"\nGenerated by Afaan Oromo STT System\n{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(150, 150, 150)
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.save(output_path)
    return output_path


def export_pdf(
    text: str,
    output_path: str,
    title: str = "Transcription",
    metadata: Optional[Dict] = None,
    segments: Optional[List[Dict]] = None
) -> str:
    """
    Export transcription as PDF file
    
    Args:
        text: Transcription text
        output_path: Output file path
        title: Document title
        metadata: Optional metadata
        segments: Optional segments with timestamps
    
    Returns:
        Path to created file
    """
    try:
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    except ImportError:
        raise ImportError("reportlab not installed. Install with: pip install reportlab")
    
    doc = SimpleDocTemplate(output_path, pagesize=A4)
    story = []
    styles = getSampleStyleSheet()
    
    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1f2937'),
        spaceAfter=30,
        alignment=TA_CENTER
    )
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 0.3*inch))
    
    # Metadata
    if metadata:
        story.append(Paragraph("Metadata", styles['Heading2']))
        meta_data = [[k.replace('_', ' ').title(), str(v)] for k, v in metadata.items()]
        meta_table = Table(meta_data, colWidths=[2*inch, 4*inch])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1f2937')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 0.3*inch))
    
    # Transcription
    story.append(Paragraph("Transcription", styles['Heading2']))
    story.append(Spacer(1, 0.2*inch))
    
    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['BodyText'],
        fontSize=11,
        leading=16,
        alignment=TA_JUSTIFY,
        spaceAfter=12
    )
    
    timestamp_style = ParagraphStyle(
        'Timestamp',
        parent=styles['BodyText'],
        fontSize=9,
        textColor=colors.HexColor('#6b7280'),
        spaceAfter=4
    )
    
    if segments:
        # Show with timestamps
        for segment in segments:
            start = segment.get('start', 0)
            end = segment.get('end', 0)
            segment_text = segment.get('text', '').strip()
            
            if not segment_text:
                continue
            
            # Timestamp
            timestamp = f"[{format_timestamp_srt(start).split(',')[0]} - {format_timestamp_srt(end).split(',')[0]}]"
            story.append(Paragraph(timestamp, timestamp_style))
            
            # Text
            story.append(Paragraph(segment_text, body_style))
    else:
        # Plain text (split into paragraphs)
        paragraphs = text.split('\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para, body_style))
    
    # Footer
    story.append(Spacer(1, 0.5*inch))
    footer_style = ParagraphStyle(
        'Footer',
        parent=styles['BodyText'],
        fontSize=8,
        textColor=colors.HexColor('#9ca3af'),
        alignment=TA_CENTER
    )
    footer_text = f"Generated by Afaan Oromo STT System<br/>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    story.append(Paragraph(footer_text, footer_style))
    
    doc.build(story)
    return output_path


def export_json(
    data: Dict,
    output_path: str,
    pretty: bool = True
) -> str:
    """
    Export transcription data as JSON
    
    Args:
        data: Data dictionary
        output_path: Output file path
        pretty: Pretty print JSON
    
    Returns:
        Path to created file
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        if pretty:
            json.dump(data, f, ensure_ascii=False, indent=2)
        else:
            json.dump(data, f, ensure_ascii=False)
    
    return output_path


def export_all_formats(
    text: str,
    segments: List[Dict],
    base_filename: str,
    output_dir: str,
    metadata: Optional[Dict] = None
) -> Dict[str, str]:
    """
    Export transcription in all available formats
    
    Args:
        text: Transcription text
        segments: Subtitle segments
        base_filename: Base filename (without extension)
        output_dir: Output directory
        metadata: Optional metadata
    
    Returns:
        Dictionary mapping format -> file path
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    results = {}
    
    # TXT
    try:
        txt_path = output_dir / f"{base_filename}.txt"
        export_txt(text, str(txt_path), metadata)
        results['txt'] = str(txt_path)
    except Exception as e:
        print(f"[ERROR] TXT export failed: {e}")
    
    # SRT
    try:
        srt_path = output_dir / f"{base_filename}.srt"
        export_srt(segments, str(srt_path))
        results['srt'] = str(srt_path)
    except Exception as e:
        print(f"[ERROR] SRT export failed: {e}")
    
    # VTT
    try:
        vtt_path = output_dir / f"{base_filename}.vtt"
        export_vtt(segments, str(vtt_path))
        results['vtt'] = str(vtt_path)
    except Exception as e:
        print(f"[ERROR] VTT export failed: {e}")
    
    # DOCX
    try:
        docx_path = output_dir / f"{base_filename}.docx"
        export_docx(text, str(docx_path), base_filename, metadata, segments)
        results['docx'] = str(docx_path)
    except Exception as e:
        print(f"[ERROR] DOCX export failed: {e}")
    
    # PDF
    try:
        pdf_path = output_dir / f"{base_filename}.pdf"
        export_pdf(text, str(pdf_path), base_filename, metadata, segments)
        results['pdf'] = str(pdf_path)
    except Exception as e:
        print(f"[ERROR] PDF export failed: {e}")
    
    # JSON
    try:
        json_path = output_dir / f"{base_filename}.json"
        json_data = {
            'text': text,
            'segments': segments,
            'metadata': metadata or {},
            'exported_at': datetime.now().isoformat()
        }
        export_json(json_data, str(json_path))
        results['json'] = str(json_path)
    except Exception as e:
        print(f"[ERROR] JSON export failed: {e}")
    
    return results
